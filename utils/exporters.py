"""
exporters.py
- Sample/template docx banata hai jisse user samajh sake format kya hai
- Question bank ko docx me export karta hai
- User scorecard text banata hai
"""

import os
import json
import time
from docx import Document
from docx.shared import Pt

EXPORT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


def make_sample_template():
    path = os.path.join(EXPORT_DIR, "sample_question_format.docx")
    doc = Document()
    doc.add_heading("Quiz Question Format Sample", level=1)
    doc.add_paragraph("Subject: Physics")
    doc.add_paragraph("Chapter: Motion")
    doc.add_paragraph("")
    doc.add_paragraph("Q1. Speed ka SI unit kya hai?")
    doc.add_paragraph("A) Kg")
    doc.add_paragraph("B) m/s")
    doc.add_paragraph("C) Newton")
    doc.add_paragraph("D) Watt")
    doc.add_paragraph("Answer: B")
    doc.add_paragraph("Explanation: Speed = distance/time hota hai, unit m/s.")
    doc.add_paragraph("")
    doc.add_paragraph("Q2. Newton ka pehla niyam kis naam se jana jata hai?")
    doc.add_paragraph("A) Jadatva ka niyam")
    doc.add_paragraph("B) Gravity ka niyam")
    doc.add_paragraph("C) Urja Samrakshan")
    doc.add_paragraph("D) Koi nahi")
    doc.add_paragraph("Answer: A")
    doc.save(path)
    return path


def export_questions_to_docx(questions, filename="question_bank_export.docx"):
    """questions: list of sqlite Row objects from database.get_questions()"""
    path = os.path.join(EXPORT_DIR, filename)
    doc = Document()
    doc.add_heading("Question Bank Export", level=1)
    for i, q in enumerate(questions, 1):
        opts = json.loads(q["options"])
        p = doc.add_paragraph()
        run = p.add_run(f"Q{i}. {q['question_text']}")
        run.bold = True
        for idx, opt in enumerate(opts):
            letter = chr(65 + idx)
            doc.add_paragraph(f"{letter}) {opt}")
        correct_letter = chr(65 + q["correct_index"])
        doc.add_paragraph(f"Answer: {correct_letter}")
        if q["explanation"]:
            doc.add_paragraph(f"Explanation: {q['explanation']}")
        doc.add_paragraph("")
    doc.save(path)
    return path


def build_scorecard_text(user_row, session_results_row, total_questions, chat_title):
    name = user_row["first_name"] or user_row["username"] or "Player"
    correct = session_results_row["correct"] if session_results_row else 0
    wrong = session_results_row["wrong"] if session_results_row else 0
    coins = session_results_row["coins_earned"] if session_results_row else 0
    attempted = correct + wrong
    accuracy = round((correct / attempted) * 100, 1) if attempted else 0

    text = (
        f"🏆 *Quiz Scorecard*\n"
        f"👤 Player: {name}\n"
        f"📍 Group: {chat_title}\n"
        f"📝 Total Questions: {total_questions}\n"
        f"✅ Correct: {correct}\n"
        f"❌ Wrong: {wrong}\n"
        f"🎯 Accuracy: {accuracy}%\n"
        f"🪙 Coins Earned: {coins}\n"
    )
    return text


def build_leaderboard_text(rows, title="🏆 Leaderboard"):
    if not rows:
        return f"{title}\n\nAbhi tak koi data nahi hai."
    lines = [f"{title}\n"]
    medals = ["🥇", "🥈", "🥉"]
    for i, r in enumerate(rows):
        prefix = medals[i] if i < 3 else f"{i+1}."
        name = r["first_name"] or r["username"] or f"User{r['user_id']}"
        lines.append(f"{prefix} {name} — 🪙 {r['coins']} coins")
    return "\n".join(lines)
